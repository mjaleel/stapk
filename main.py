import flet as ft
import openpyxl
import re

def check_iban_mod97(iban):
    try:
        iban = str(iban).replace(" ", "").strip().upper()
        if not iban.startswith("IQ") or len(iban) != 23:
            return False
        rearranged = iban[4:] + iban[:4]
        numeric_iban = ""
        for char in rearranged:
            if char.isdigit():
                numeric_iban += char
            else:
                numeric_iban += str(ord(char) - 55)
        return int(numeric_iban) % 97 == 1
    except:
        return False

def clean_amount_val(val):
    if val is None or str(val).strip() == "" or str(val).lower() == "none":
        return 0.0
    val_str = str(val).strip()
    arabic_map = {'٠':'0','١':'1','٢':'2','٣':'3','٤':'4','٥':'5','٦':'6','٧':'7','٨':'8','٩':'9'}
    for arb, eng in arabic_map.items():
        val_str = val_str.replace(arb, eng)
    val_str = val_str.replace(",", "")
    if val_str.count(".") > 1:
        val_str = val_str.replace(".", "")
    val_str = re.sub(r'[^\d.]', '', val_str)
    try:
        return float(val_str)
    except:
        return 0.0

def main(page: ft.Page):
    page.title = "مدقق الرواتب الذكي"
    page.rtl = True
    page.scroll = ft.ScrollMode.AUTO
    page.theme_mode = ft.ThemeMode.LIGHT
    page.padding = 20

    state = {"wb": None, "file_path": ""}

    title = ft.Text("🏦 نظام تدقيق ومعالجة الرواتب", size=22, weight=ft.FontWeight.BOLD)
    subtitle = ft.Text("ارفع ملف الإكسل ثم اختر العملية المطلوبة.", size=14, color="grey")
    file_status = ft.Text("", color="blue", weight=ft.FontWeight.BOLD)

    sheet_dd = ft.Dropdown(label="📄 اختر ورقة العمل", visible=False)
    iban_dd = ft.Dropdown(label="🎯 اختر عمود الـ IBAN", visible=False)
    amount_dd = ft.Dropdown(label="💰 اختر عمود المبلغ", visible=False)
    
    progress_ring = ft.ProgressBar(visible=False)
    results_col = ft.Column(spacing=8)

    def load_columns(e):
        sheet_name = sheet_dd.value
        ws = state["wb"][sheet_name]
        headers = [str(cell.value or f"عمود {i+1}") for i, cell in enumerate(ws[1])]
        
        iban_dd.options = [ft.dropdown.Option(h) for h in headers]
        iban_dd.value = headers[0] if headers else None
        iban_dd.visible = True

        amount_options = ["تخطي فحص المبلغ"] + headers
        amount_dd.options = [ft.dropdown.Option(h) for h in amount_options]
        amount_dd.value = "تخطي فحص المبلغ"
        amount_dd.visible = True

        audit_btn.visible = True
        clean_btn.visible = True
        results_col.controls.clear()
        page.update()

    sheet_dd.on_change = load_columns

    def on_file_picked(e: ft.FilePickerResultEvent):
        if e.files:
            file_path = e.files[0].path
            state["file_path"] = file_path
            file_status.value = f"📂 الملف المحدد: {e.files[0].name}"
            
            state["wb"] = openpyxl.load_workbook(file_path, data_only=True)
            sheet_dd.options = [ft.dropdown.Option(s) for s in state["wb"].sheetnames]
            sheet_dd.value = state["wb"].sheetnames[0]
            sheet_dd.visible = True
            load_columns(None)

    def on_file_saved(e: ft.FilePickerResultEvent):
        if e.path:
            try:
                state["wb"].save(e.path)
                results_col.controls.clear()
                results_col.controls.append(
                    ft.Container(
                        content=ft.Text(f"✅ تم حفظ الملف النظيف بنجاح في:\n{e.path}", color="green", weight=ft.FontWeight.BOLD),
                        padding=10, bgcolor="#f0fff4", border_radius=8
                    )
                )
            except Exception as ex:
                results_col.controls.append(ft.Text(f"❌ خطأ أثناء الحفظ: {str(ex)}", color="red"))
            page.update()

    pick_dialog = ft.FilePicker(on_result=on_file_picked)
    save_dialog = ft.FilePicker(on_result=on_file_saved)
    page.overlay.extend([pick_dialog, save_dialog])

    upload_btn = ft.ElevatedButton(
        "📂 اختر ملف الرواتب (Excel)",
        on_click=lambda _: pick_dialog.pick_files(allowed_extensions=["xlsx"]),
        bgcolor="#3182ce", color="white", height=45
    )

    def run_audit(e):
        results_col.controls.clear()
        progress_ring.visible = True
        page.update()

        state["wb"] = openpyxl.load_workbook(state["file_path"], data_only=True)
        ws = state["wb"][sheet_dd.value]
        headers = [str(cell.value or f"عمود {i+1}") for i, cell in enumerate(ws[1])]
        
        iban_idx = headers.index(iban_dd.value)
        amt_idx = headers.index(amount_dd.value) if amount_dd.value != "تخطي فحص المبلغ" else -1

        critical_errors = []
        warnings = []
        seen_ibans = {}

        for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            raw_iban_val = str(row[iban_idx] if row[iban_idx] is not None else "")
            
            if raw_iban_val.strip() == "" or raw_iban_val.lower() == "none":
                critical_errors.append(f"❌ [صف {row_num}] خطأ: حقل الآيبان فارغ!")
                continue
            if " " in raw_iban_val:
                warnings.append(f"⚠️ [صف {row_num}] تنبيه: توجد مسافات في الآيبان: {raw_iban_val}")
            if re.search(r'[a-z]', raw_iban_val):
                critical_errors.append(f"❌ [صف {row_num}] خطأ: أحرف صغيرة: {raw_iban_val}")
            
            clean_iban = raw_iban_val.replace(" ", "").upper()
            if not check_iban_mod97(clean_iban):
                critical_errors.append(f"❌ [صف {row_num}] خطأ بنكي: الحساب غير صحيح: {clean_iban}")
            
            if clean_iban in seen_ibans:
                warnings.append(f"📝 [صف {row_num}] مكرر: مع الصف {seen_ibans[clean_iban]}")
            else:
                seen_ibans[clean_iban] = row_num

            if amt_idx != -1:
                raw_amt_val = str(row[amt_idx] if row[amt_idx] is not None else "")
                amt = clean_amount_val(raw_amt_val)
                if raw_amt_val.strip() == "" or raw_amt_val.lower() == "none":
                    critical_errors.append(f"❌ [صف {row_num}] خطأ: المبلغ فارغ!")
                elif amt <= 0:
                    critical_errors.append(f"❌ [صف {row_num}] خطأ: المبلغ غير صالح: {raw_amt_val}")

        progress_ring.visible = False

        if not critical_errors and not warnings:
            results_col.controls.append(
                ft.Text("✅ الملف سليم تماماً وخالٍ من الأخطاء.", color="green", size=18, weight=ft.FontWeight.BOLD)
            )
        else:
            if critical_errors:
                results_col.controls.append(
                    ft.Text(f"⛔ تم اكتشاف {len(critical_errors)} أخطاء حرجة:", color="red", size=16, weight=ft.FontWeight.BOLD)
                )
                for err in critical_errors:
                    results_col.controls.append(ft.Text(err, color="red", size=13))
            if warnings:
                results_col.controls.append(
                    ft.Text(f"📝 التنبيهات ({len(warnings)}):", color="orange", size=16, weight=ft.FontWeight.BOLD)
                )
                for w in warnings:
                    results_col.controls.append(ft.Text(w, color="#dd6b20", size=13))
        page.update()

    def clean_and_save(e):
        results_col.controls.clear()
        progress_ring.visible = True
        page.update()

        state["wb"] = openpyxl.load_workbook(state["file_path"], data_only=True)
        ws = state["wb"][sheet_dd.value]
        headers = [str(cell.value or f"عمود {i+1}") for i, cell in enumerate(ws[1])]
        
        iban_idx = headers.index(iban_dd.value) + 1
        amt_idx = headers.index(amount_dd.value) + 1 if amount_dd.value != "تخطي فحص المبلغ" else -1

        for row_num in range(2, ws.max_row + 1):
            iban_cell = ws.cell(row=row_num, column=iban_idx)
            if iban_cell.value is not None:
                iban_cell.value = str(iban_cell.value).replace(" ", "").strip().upper()
            
            if amt_idx != -1:
                amt_cell = ws.cell(row=row_num, column=amt_idx)
                amt_cell.value = f"{clean_amount_val(amt_cell.value):.0f}"

        progress_ring.visible = False
        page.update()

        save_dialog.save_file(
            dialog_title="حفظ ملف الإكسل النظيف",
            file_name=f"Cleaned_{sheet_dd.value}.xlsx",
            allowed_extensions=["xlsx"]
        )

    audit_btn = ft.ElevatedButton("🔍 ابدأ التدقيق", on_click=run_audit, visible=False, bgcolor="#2b6cb0", color="white")
    clean_btn = ft.ElevatedButton("⚙️ معالجة وحفظ الملف", on_click=clean_and_save, visible=False, bgcolor="#38a169", color="white")

    page.add(
        title,
        subtitle,
        ft.Divider(),
        upload_btn,
        file_status,
        sheet_dd,
        iban_dd,
        amount_dd,
        ft.Row([audit_btn, clean_btn], wrap=True, spacing=10),
        progress_ring,
        ft.Divider(),
        results_col
    )

ft.app(target=main)